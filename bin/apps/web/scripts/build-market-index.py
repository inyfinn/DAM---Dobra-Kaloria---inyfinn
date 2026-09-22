# -*- coding: utf-8 -*-
"""Buduje data/market-index.json z teczek projektow (etap) i kart zalozenia indeksow (GTIN).

Zrodla (poza repo, na dysku uzytkownika):
  <ROOT>/<rok>/Projekty w toku|wdrozone|zawieszone/.../<projekt>/...
      -> etap projektu; indeksy opakowan 6xxxxxx wyciagane z nazw plikow i folderow
  <ROOT>/Karty zalozenia indeksow/*.xlsx
      -> indeks handlowy 33xxxxx, GTIN jednostkowy/zbiorczy, klient, marka

Wynik: mapa indeks_opakowania -> {etap, projekt, rok, marka, indeks handlowy, GTIN}.
Indeks opakowania 6xxxxxx to ten sam numer, ktorym nazwane sa foldery rewizji w DAM,
wiec karta projektu moze pokazac, czy dane opakowanie jest juz w obrocie.

Gdy ROOT nie istnieje (inna maszyna, brak dysku), skrypt konczy sie kodem 0
i zostawia poprzedni plik nietkniety - brak teczek nie jest bledem.
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

WEB = Path(__file__).resolve().parents[1]
OUT_FILE = WEB / "data" / "market-index.json"

DEFAULT_ROOT = os.environ.get("DAM_TECZKI_ROOT", r"D:\Teczki projektow")
# Nazwa katalogu bywa zapisana z polskimi znakami - obie formy sa akceptowane.
ROOT_CANDIDATES = [
    "D:\\Teczki projekt\u00f3w",
    "D:\\Teczki projektow",
    "X:\\Teczki projekt\u00f3w",
    "X:\\Teczki projektow",
]

# Indeks opakowania: 7 cyfr zaczynajacych sie od 6 (6300xxx folie/rekawy, 6900xxx etykiety).
PACK_INDEX_RE = re.compile(r"(?<!\d)(6\d{6})(?!\d)")
# Indeks handlowy towaru: 7 cyfr zaczynajacych sie od 33.
TRADE_INDEX_RE = re.compile(r"(?<!\d)(33\d{5})(?!\d)")
# Folder projektu zaczyna sie od roku: "2026_(DK) Roslinna kofta", "2025 (DK) Owsianka".
PROJECT_DIR_RE = re.compile(r"^(20\d{2})[_\s]")
# Sufiksy konfliktow Synology - nigdy nie sa zrodlem indeksu.
CONFLICT_RE = re.compile(r"_INYFINN_|_Conflict|conflict_current", re.IGNORECASE)

STAGE_BY_DIR = {
    "projekty w toku": "w_toku",
    "projekty wdrozone": "wdrozony",
    "projekty zawieszone": "zawieszony",
}
STAGE_LABELS = {
    "w_toku": {"pl": "Projekt w toku", "en": "In progress"},
    "wdrozony": {"pl": "Wdrozony - w obrocie", "en": "Implemented - on market"},
    "zawieszony": {"pl": "Zawieszony", "en": "Suspended"},
}
# Kolejnosc pewnosci: wdrozony wygrywa z w toku, w toku z zawieszonym.
STAGE_RANK = {"wdrozony": 3, "w_toku": 2, "zawieszony": 1}

DIACRITICS = str.maketrans(
    "\u0105\u0107\u0119\u0142\u0144\u00f3\u015b\u017a\u017c\u0104\u0106\u0118\u0141\u0143\u00d3\u015a\u0179\u017b",
    "acelnoszzACELNOSZZ",
)


def fold(text: str) -> str:
    return text.translate(DIACRITICS).strip().lower()


def resolve_root(explicit: str | None) -> Path | None:
    cands = [explicit] if explicit else []
    cands.append(DEFAULT_ROOT)
    cands.extend(ROOT_CANDIDATES)
    for c in cands:
        if not c:
            continue
        p = Path(c)
        if p.is_dir():
            return p
    return None


def indexes_from_name(name: str) -> list[str]:
    """Indeksy opakowan z nazwy pliku lub folderu, z pominieciem konfliktow Synology."""
    if CONFLICT_RE.search(name):
        return []
    return PACK_INDEX_RE.findall(name)


def scan_projects(root: Path) -> tuple[dict, list]:
    """Zwraca (mapa indeks -> wpis projektu, lista projektow)."""
    packages: dict[str, dict] = {}
    projects: list[dict] = []
    for year_dir in sorted(root.iterdir()):
        if not year_dir.is_dir() or not re.fullmatch(r"20\d{2}", year_dir.name):
            continue
        for status_dir in sorted(year_dir.iterdir()):
            if not status_dir.is_dir():
                continue
            stage = STAGE_BY_DIR.get(fold(status_dir.name))
            if not stage:
                continue
            for proj_dir, brand in find_project_dirs(status_dir):
                rec = collect_project(proj_dir, brand, stage, year_dir.name, root)
                projects.append(rec)
                for idx in rec["pack_indexes"]:
                    prev = packages.get(idx)
                    if prev is None or STAGE_RANK.get(stage, 0) > STAGE_RANK.get(prev["stage"], 0):
                        packages[idx] = {
                            "stage": stage,
                            "project": rec["project"],
                            "year": rec["year"],
                            "brand": rec["brand"],
                            "source": "teczki",
                        }
    return packages, projects


def find_project_dirs(status_dir: Path):
    """Foldery projektow leza na roznej glebokosci - rozpoznajemy je po prefiksie roku."""
    out = []
    for cur, dirs, _files in os.walk(status_dir):
        curp = Path(cur)
        for d in list(dirs):
            if PROJECT_DIR_RE.match(d):
                rel = curp.relative_to(status_dir)
                brand = rel.parts[0] if rel.parts else ""
                out.append((curp / d, brand))
                dirs.remove(d)  # nie schodz glebiej, projekt jest lisciem
    return out


def collect_project(proj_dir: Path, brand: str, stage: str, year: str, root: Path) -> dict:
    found: set[str] = set()
    trade: set[str] = set()
    for name in [proj_dir.name]:
        found.update(indexes_from_name(name))
        trade.update(TRADE_INDEX_RE.findall(name))
    for cur, dirs, files in os.walk(proj_dir):
        for name in list(dirs) + list(files):
            found.update(indexes_from_name(name))
            if not CONFLICT_RE.search(name):
                trade.update(TRADE_INDEX_RE.findall(name))
    return {
        "project": PROJECT_DIR_RE.sub("", proj_dir.name).strip(" _-"),
        "folder": proj_dir.name,
        "year": year,
        "brand": brand,
        "stage": stage,
        "pack_indexes": sorted(found),
        "trade_indexes": sorted(trade),
        "rel_path": str(proj_dir.relative_to(root)),
    }


def scan_cards(root: Path) -> dict:
    """Karty zalozenia indeksow: indeks opakowania -> indeks handlowy + GTIN."""
    cards_dir = None
    for d in root.iterdir():
        if d.is_dir() and fold(d.name).startswith("karty zalozenia"):
            cards_dir = d
            break
    if cards_dir is None:
        return {}
    try:
        import openpyxl  # noqa: PLC0415
    except ImportError:
        sys.stderr.write("openpyxl niedostepny - pomijam karty zalozenia indeksow\n")
        return {}

    out: dict[str, dict] = {}
    for f in sorted(cards_dir.glob("*.xlsx")):
        if f.name.upper().startswith("AAA_WZ"):
            continue
        try:
            wb = openpyxl.load_workbook(f, data_only=True)
        except Exception as exc:  # plik otwarty w Excelu, uszkodzony itp.
            sys.stderr.write(f"pomijam {f.name}: {exc}\n")
            continue
        ws = wb[wb.sheetnames[0]]
        card = {"trade_index": None, "gtin_unit": None, "gtin_bulk": None,
                "product": None, "brand": None, "client": None, "card_file": f.name}
        packs: set[str] = set()
        for r in range(1, ws.max_row + 1):
            lab = ws.cell(r, 1).value
            lab = re.sub(r"\s+", " ", str(lab)).strip().lower() if lab else ""
            val = ws.cell(r, 2).value
            sval = str(val).strip() if val is not None else None
            if lab.startswith("nr indeksu"):
                card["trade_index"] = sval
            elif lab.startswith("nazwa produktu"):
                card["product"] = sval
            elif lab.startswith("nazwa marki"):
                card["brand"] = sval
            elif lab.startswith("gtin jednostkowy"):
                card["gtin_unit"] = sval
            elif lab.startswith("gtin zbiorczy"):
                card["gtin_bulk"] = sval
            elif lab.startswith("klient"):
                card["client"] = sval
            for c in range(1, ws.max_column + 1):
                cv = ws.cell(r, c).value
                if cv is not None:
                    packs.update(PACK_INDEX_RE.findall(str(cv)))
        for idx in packs:
            out.setdefault(idx, dict(card))
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--root", help="katalog Teczki projektow")
    ap.add_argument("--out", help="plik wynikowy")
    ap.add_argument("--quiet", action="store_true")
    args = ap.parse_args()

    root = resolve_root(args.root)
    out_file = Path(args.out) if args.out else OUT_FILE
    if root is None:
        if not args.quiet:
            sys.stderr.write("Teczki projektow niedostepne - zostawiam poprzedni market-index.json\n")
        return 0

    packages, projects = scan_projects(root)
    cards = scan_cards(root)

    for idx, card in cards.items():
        entry = packages.setdefault(idx, {"stage": None, "project": None, "year": None,
                                          "brand": None, "source": "karta"})
        entry["trade_index"] = card.get("trade_index")
        entry["gtin_unit"] = card.get("gtin_unit")
        entry["gtin_bulk"] = card.get("gtin_bulk")
        entry["client"] = card.get("client")
        entry["card_file"] = card.get("card_file")
        if not entry.get("project"):
            entry["project"] = card.get("product")

    doc = {
        "version": 1,
        "updated_at": datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds"),
        "source_root": str(root),
        "stage_labels": STAGE_LABELS,
        "counts": {
            "projects": len(projects),
            "packages": len(packages),
            "with_gtin": sum(1 for v in packages.values() if v.get("gtin_unit")),
        },
        "packages": dict(sorted(packages.items())),
        "projects": projects,
    }
    out_file.parent.mkdir(parents=True, exist_ok=True)
    out_file.write_text(json.dumps(doc, ensure_ascii=False, indent=1), encoding="utf-8")
    if not args.quiet:
        print(f"projekty: {len(projects)}  opakowania: {len(packages)}  z GTIN: {doc['counts']['with_gtin']}")
        print(f"zapisano: {out_file}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
